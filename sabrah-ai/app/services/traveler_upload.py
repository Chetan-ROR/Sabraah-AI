"""Parse traveler lists from uploaded spreadsheets."""

from __future__ import annotations

import csv
import io
from typing import Any


def parse_traveler_file(content: bytes, filename: str) -> list[dict[str, Any]]:
    name = (filename or "").lower()
    if name.endswith(".xlsx") or name.endswith(".xlsm"):
        return _parse_xlsx(content)
    if name.endswith(".csv"):
        return _parse_csv(content)
    raise ValueError("Please upload a .xlsx or .csv file with traveler names.")


def _parse_csv(content: bytes) -> list[dict[str, Any]]:
    text = content.decode("utf-8-sig", errors="replace")
    reader = csv.DictReader(io.StringIO(text))
    return _rows_from_mapping(reader)


def _parse_xlsx(content: bytes) -> list[dict[str, Any]]:
    try:
        from openpyxl import load_workbook
    except ImportError as exc:
        raise ValueError(
            "Excel support is not installed on the server. Use CSV instead."
        ) from exc

    wb = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
    sheet = wb.active
    rows = list(sheet.iter_rows(values_only=True))
    if not rows:
        return []
    headers = [_norm_header(h) for h in rows[0]]
    mapped = []
    for raw in rows[1:]:
        if not raw or not any(raw):
            continue
        row = {headers[i]: raw[i] for i in range(min(len(headers), len(raw)))}
        mapped.append(row)
    return _rows_from_mapping(mapped)


def _norm_header(value: Any) -> str:
    text = str(value or "").strip().lower().replace(" ", "_")
    aliases = {
        "passenger_name": "name",
        "full_name": "name",
        "traveler_name": "name",
        "traveler": "name",
        "mobile": "phone",
        "contact": "phone",
        "phone_number": "phone",
        "meal_preference": "meal",
        "diet": "meal",
        "allergy": "allergies",
    }
    return aliases.get(text, text)


def _rows_from_mapping(rows: Any) -> list[dict[str, Any]]:
    out: list[dict[str, Any]] = []
    for row in rows:
        if not isinstance(row, dict):
            continue
        norm = {_norm_header(k): v for k, v in row.items()}
        name = _cell(norm.get("name"))
        if not name:
            continue
        out.append(
            {
                "name": name,
                "phone": _cell(norm.get("phone")),
                "age": _cell(norm.get("age")),
                "meal": _cell(norm.get("meal")),
                "allergies": _cell(norm.get("allergies")),
                "seat_preference": _cell(norm.get("seat_preference")),
            }
        )
    return out


def _cell(value: Any) -> str:
    if value is None:
        return ""
    return str(value).strip()
